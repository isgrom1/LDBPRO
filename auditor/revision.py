"""Copia de revisión: celdas coloreadas y comentadas + hojas RESUMEN y HALLAZGOS.

Se escribe con openpyxl. openpyxl pierde las listas desplegables al guardar; se acepta
porque este archivo es solo para mirar. El archivo corregido se escribe aparte, en XML.
"""

from __future__ import annotations

import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .hallazgos import REGLAS, Hallazgo, Severidad
from .motor import Puntaje, formula_puntaje
from .texto import mostrar

AUTOR = "Auditor XPE"
_NEGRITA = Font(bold=True)
_ENCABEZADO = PatternFill("solid", fgColor="FF305496")
_FUENTE_ENC = Font(bold=True, color="FFFFFFFF")
_AJUSTE = Alignment(wrap_text=True, vertical="top")


def _relleno(sev: Severidad) -> PatternFill:
    return PatternFill("solid", fgColor=sev.color)


def _celda_editable(ws, ref: str):
    c = ws[ref]
    if isinstance(c, MergedCell):
        for rango in ws.merged_cells.ranges:
            if ref in rango:
                return ws.cell(rango.min_row, rango.min_col)
    return c


def _texto_comentario(h: Hallazgo, aplicado: bool, relacionado: str | None = None) -> str:
    cab = f"[{h.id_regla} · {h.severidad.etiqueta}]"
    if relacionado:
        cab += f" (relacionado con {relacionado})"
    estado = "\n✔ Corregido automáticamente en _CORREGIDA." if aplicado else ""
    return f"{cab}\n{h.problema}\n→ {h.sugerencia}{estado}"


def marcar_celdas(wb, hallazgos: list[Hallazgo], aplicados: set[tuple[str, str]]) -> int:
    por_celda: dict[tuple[str, str], list[tuple[Hallazgo, str | None]]] = defaultdict(list)
    for h in hallazgos:
        if h.hoja not in wb.sheetnames or not h.celda:
            continue
        por_celda[(h.hoja, h.celda)].append((h, None))
        for extra in h.celdas_extra:
            por_celda[(h.hoja, extra)].append((h, f"{h.hoja}!{h.celda}"))
    for (hoja, ref), items in por_celda.items():
        ws = wb[hoja]
        c = _celda_editable(ws, ref)
        peor = max((h.severidad for h, _ in items), key=lambda s: s.nivel)
        c.fill = _relleno(peor)
        partes = [_texto_comentario(h, (h.hoja, h.celda) in aplicados and h.auto_corregible, rel) for h, rel in items]
        texto = "\n\n".join(partes)
        if c.comment is not None and c.comment.text:
            texto = f"{c.comment.text}\n\n— Auditoría —\n{texto}"
        comentario = Comment(texto[:30000], AUTOR)
        lineas = sum(1 + len(p) // 60 for p in texto.split("\n"))
        comentario.width = 420
        comentario.height = min(700, 30 + 15 * lineas)
        c.comment = comentario
    return len(por_celda)


def _escribir_fila(ws, fila: int, valores, negrita=False, relleno=None, fuente=None):
    for j, v in enumerate(valores, start=1):
        c = ws.cell(fila, j, v)
        c.alignment = _AJUSTE
        if negrita:
            c.font = _NEGRITA
        if fuente:
            c.font = fuente
        if relleno:
            c.fill = relleno


def hoja_hallazgos(wb, hallazgos: list[Hallazgo], aplicados: set, pendientes: set) -> None:
    ws = wb.create_sheet("HALLAZGOS", 0)
    columnas = [
        ("N°", 6), ("Severidad", 16), ("Regla", 7), ("Descripción de la regla", 34), ("Hoja", 13),
        ("Celda", 9), ("Fila", 6), ("ID Estación", 11), ("Valor actual", 22), ("Problema", 60),
        ("Sugerencia", 50), ("Auto-corregible", 11), ("Aplicado en _CORREGIDA", 12),
        ("Valor nuevo", 18), ("Sigue pendiente", 12), ("Decisión / Corrección", 40),
    ]
    _escribir_fila(ws, 1, [c for c, _ in columnas], relleno=_ENCABEZADO, fuente=_FUENTE_ENC)
    for j, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    for i, h in enumerate(hallazgos, start=1):
        r = REGLAS.get(h.id_regla)
        aplicado = h.auto_corregible and (h.hoja, h.celda) in aplicados
        fila = [
            i, h.severidad.etiqueta, h.id_regla, r.titulo if r else "", h.hoja, h.celda or "",
            h.fila or "", h.id_estacion or "", mostrar(h.valor_actual) if h.valor_actual is not None else "",
            h.problema, h.sugerencia, "Sí" if h.auto_corregible else "No", "Sí" if aplicado else "No",
            mostrar(h.correccion.valor_nuevo) if h.correccion else "",
            "Sí" if h.firma() in pendientes else "No", "",
        ]
        _escribir_fila(ws, i + 1, fila)
        ws.cell(i + 1, 2).fill = _relleno(h.severidad)
        if h.celda and h.hoja in wb.sheetnames:
            c = ws.cell(i + 1, 6)
            c.hyperlink = f"#'{h.hoja}'!{h.celda}"
            c.style = "Hyperlink"
            c.alignment = _AJUSTE  # el estilo Hyperlink trae alineación inferior
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(hallazgos) + 1}"


def hoja_resumen(
    wb, nombre: str, formato: str, antes: Puntaje, despues: Puntaje, hallazgos_antes, hallazgos_despues,
    config: dict, estadisticas: dict, n_cambios: int,
) -> None:
    ws = wb.create_sheet("RESUMEN", 0)
    # A: indicador / regla · B-C: antes-después o severidad-n.º · D-E: textos largos
    for col, ancho in zip("ABCDE", (52, 28, 28, 60, 44)):
        ws.column_dimensions[col].width = ancho
    ws["A1"] = "Auditoría de planilla de fauna XPE"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"Archivo: {nombre} · formato real: {formato}"
    fila = 4
    _escribir_fila(ws, fila, ["Indicador", "Antes", f"Después de {n_cambios} auto-correcciones"],
                   relleno=_ENCABEZADO, fuente=_FUENTE_ENC)
    filas = [
        ("Estado", antes.estado, despues.estado),
        ("Índice de calidad global (0–100)", round(antes.indice, 1), round(despues.indice, 1)),
        ("% registros limpios", f"{antes.pct_registros:.1f}% ({antes.n_registros_limpios}/{antes.n_registros})",
         f"{despues.pct_registros:.1f}% ({despues.n_registros_limpios}/{despues.n_registros})"),
        ("% estaciones limpias", f"{antes.pct_estaciones:.1f}% ({antes.n_estaciones_limpias}/{antes.n_estaciones})",
         f"{despues.pct_estaciones:.1f}% ({despues.n_estaciones_limpias}/{despues.n_estaciones})"),
        ("Reglas 🔴 distintas", ", ".join(antes.reglas_criticas) or "—", ", ".join(despues.reglas_criticas) or "—"),
    ]
    filas += [(f"Hallazgos {s.etiqueta}", antes.por_severidad[s.etiqueta], despues.por_severidad[s.etiqueta]) for s in Severidad]
    for etiqueta, a, d in filas:
        fila += 1
        _escribir_fila(ws, fila, [etiqueta, a, d])
        ws.cell(fila, 1).font = _NEGRITA
    fila += 1
    notas = [
        ("Fórmula del índice", formula_puntaje(config)),
        ("Estado", "✅ LISTA = sin 🔴 ni 🟠 · ⚠️ USABLE CON REVISIÓN = sin 🔴 pero con 🟠 · ❌ NO USAR = hay 🔴"),
    ]
    notas += [
        (f"Aves en transecto entre {st['rango']} ({hoja})", f"{st['pct']:.0f}% ({st['n_en_rango']} de {st['n']})")
        for hoja, st in estadisticas.get("R11", {}).items()
    ]
    for etiqueta, texto in notas:
        fila += 1
        _escribir_fila(ws, fila, [etiqueta, texto])
        ws.cell(fila, 1).font = _NEGRITA
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        ws.row_dimensions[fila].height = 15 * (1 + len(texto) // 110)

    fila += 2
    _escribir_fila(ws, fila, ["Regla", "Severidad", "Auto-corregible", "Antes", "Después"],
                   relleno=_ENCABEZADO, fuente=_FUENTE_ENC)
    n_antes = Counter(h.id_regla for h in hallazgos_antes)
    n_despues = Counter(h.id_regla for h in hallazgos_despues)
    for r in sorted(REGLAS.values(), key=lambda r: r.orden):
        if not (n_antes[r.id] or n_despues[r.id]):
            continue
        fila += 1
        _escribir_fila(ws, fila, [f"{r.id} · {r.titulo}", r.severidad.etiqueta, "Sí" if r.auto_corregible else "No",
                                  n_antes[r.id], n_despues[r.id]])
        ws.cell(fila, 2).fill = _relleno(r.severidad)

    fila += 2
    ws.cell(fila, 1, "Top 10 de problemas pendientes (después de las auto-correcciones)").font = Font(bold=True, size=12)
    fila += 1
    _escribir_fila(ws, fila, ["Regla", "Severidad", "N.º de hallazgos", "Ejemplo", "Celdas"],
                   relleno=_ENCABEZADO, fuente=_FUENTE_ENC)
    for r_id, items in top_problemas(hallazgos_despues, 10):
        r = REGLAS.get(r_id)
        fila += 1
        celdas = ", ".join(f"{h.hoja}!{h.celda}" for h in items[:6] if h.celda) + (" …" if len(items) > 6 else "")
        _escribir_fila(ws, fila, [f"{r_id} · {r.titulo}" if r else r_id, items[0].severidad.etiqueta, len(items),
                                  items[0].problema, celdas])
        ws.cell(fila, 2).fill = _relleno(items[0].severidad)
    ws.sheet_view.showGridLines = False


def top_problemas(hallazgos: list[Hallazgo], n: int):
    por_regla = defaultdict(list)
    for h in hallazgos:
        por_regla[h.id_regla].append(h)
    ranking = sorted(por_regla.items(), key=lambda kv: (-kv[1][0].severidad.nivel, -len(kv[1]), kv[0]))
    return ranking[:n]


def escribir_revision(
    ruta_xlsx: Path, destino: Path, nombre: str, formato: str, hallazgos_antes: list[Hallazgo],
    hallazgos_despues: list[Hallazgo], antes: Puntaje, despues: Puntaje, aplicados: set,
    config: dict, estadisticas: dict, n_cambios: int,
) -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(ruta_xlsx)
    marcar_celdas(wb, hallazgos_antes, aplicados)
    pendientes = {h.firma() for h in hallazgos_despues}
    hoja_hallazgos(wb, hallazgos_antes, aplicados, pendientes)
    hoja_resumen(wb, nombre, formato, antes, despues, hallazgos_antes, hallazgos_despues, config,
                 estadisticas, n_cambios)
    wb.active = 0
    for ws in wb.worksheets[1:]:
        ws.sheet_view.tabSelected = False
    wb.save(destino)
