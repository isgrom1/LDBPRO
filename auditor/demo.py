"""Planilla de ejemplo con el formato XPE y errores típicos (datos ficticios).

Sirve para probar la aplicación sin una planilla real, para la autoprueba del ejecutable y
para las pruebas automáticas que no pueden usar datos de proyectos.
"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NARANJO = PatternFill("solid", fgColor="FFF4B183")
BLANCO = PatternFill("solid", fgColor="FFFFFFFF")
NEGRITA = Font(bold=True)
NBSP = " "
FILAS_PLANTILLA = 120  # filas de Registros con fórmulas (como la plantilla real)

ENC_ESTACIONES = [
    "N", "Proyecto", "Región", "Temporada", "ID Temporada DwC (NO LLENAR)", "Ambiente", "Ladera de Exposición",
    "ID Estación", "ID Final Estación", "ID Estación DwC (NO LLENAR)", "UTM E Estación (X)", "UTM S Estación (Y)",
    "UTM E inicio (X)", "UTM S Inicio (Y)", "UTM E fin (X)", "UTM S fin (Y)", "Transecto general", "Playback Aves",
    "Playback Anfibios", "Cámara Trampa", "Trampas Sherman",
    "Muestreado por (Nombre Apellido persona 1;Nombre Apellido persona 2; etc)",
    "Identificado por (Nombre Apellido persona 1;Nombre Apellido persona 2; etc)",
]
ENC_REGISTROS = [
    "N°", "Proyecto", "Región", "Fecha", "Año", "Mes", "Día", "Temporada", "Ambiente", "Ladera de Exposición",
    "Clima", "ID Estación", "ID Final Estación", "UTM E Estación (X)", "UTM S Estación (Y)", "UTM E inicio (X)",
    "UTM S Inicio (Y)", "UTM E fin (X)", "UTM S fin (Y)", "Transecto general", "Playback Aves",
    "Playback Anfibios", "Cámara Trampa", "Trampas Sherman", "AM/PM", "Hora", "Metodología usada para registro",
    "Tipo registro: Directo/Indirecto", "Registro: Individuo, feca, vocalización, etc", "Nombre común",
    "Nombre científico", "Estado desarrollo", "Estado del organismo", "Sexo", "Abundancia", "Comportamiento",
    "Fotos", "Observaciones", "UTM E Captura (X)", "UTM S Captura (Y)", "Reino", "Filo", "Clase", "Orden",
    "Familia", "Género", "Epíteto específico", "Epíteto infraespecífico",
    "Muestreado por (Nombre Apellido persona 1;Nombre Apellido persona 2; etc)",
    "Identificado por (Nombre Apellido persona 1;Nombre Apellido persona 2; etc)",
]
ENC_TA = [
    "N°", "Proyecto", "Región", "Fecha", "Año", "Mes", "Día", "Temporada", "Ambiente", "Ladera de Exposición",
    "Clima", "ID Estación", "ID Estación final", "UTM E Estación (X)", "UTM S Estación (Y)", "AM/PM", "Hora",
    "Nombre común", "Nombre científico", "Estado desarrollo", "Sexo", "Abundancia", "Origen", "Destino",
    "Altura vuelo", "Altura vuelo (LAT o metros)", "Comportamiento", "Fotos", "Observaciones",
]
# Columnas de Registros que son fórmula: letra -> columna de Estaciones que traen (INDEX/MATCH por ID)
DESDE_ESTACIONES = {
    "B": "B", "C": "C", "H": "D", "I": "F", "J": "G", "M": "I", "N": "K", "O": "L", "P": "M", "Q": "N",
    "R": "O", "S": "P", "T": "Q", "U": "R", "V": "S", "W": "T", "X": "U", "AW": "V", "AX": "W",
}
DESDE_BD = {"AE": "B", "AO": "C", "AP": "D", "AQ": "E", "AR": "F", "AS": "G", "AT": "H", "AU": "I", "AV": "J"}
FORMULA_DIR_IND = (
    'IF(OR(AC{r}="Individuo",AC{r}="Vocalización"),"Directo",IF(OR(AC{r}="Fecas",AC{r}="Madriguera",'
    'AC{r}="Cururera",AC{r}="Plumas",AC{r}="Muda",AC{r}="Huellas",AC{r}="Huesos",AC{r}="Nido",'
    'AC{r}="Registro de audio",AC{r}="Egagrópila"),"Indirecto","REVISAR"))'
)

# Nombre común, científico, clase, orden, familia
ESPECIES = [
    ("Chincol", "Zonotrichia capensis", "Aves", "Passeriformes", "Passerellidae"),
    ("Loica", "Leistes loyca", "Aves", "Passeriformes", "Icteridae"),
    ("Tenca", "Mimus thenca", "Aves", "Passeriformes", "Mimidae"),
    ("Diuca", "Diuca diuca", "Aves", "Passeriformes", "Thraupidae"),
    ("Cometocino de Gay", "Phrygilus gayi", "Aves", "Passeriformes", "Thraupidae"),
    ("Perdiz chilena", "Nothoprocta perdicaria", "Aves", "Tinamiformes", "Tinamidae"),
    ("Picaflor chico", "Sephanoides sephaniodes", "Aves", "Apodiformes", "Trochilidae"),
    ("Turca", "Pteroptochos megapodius", "Aves", "Passeriformes", "Rhinocryptidae"),
    ("Chercán", "Troglodytes aedon", "Aves", "Passeriformes", "Troglodytidae"),
    ("Jilguero", "Spinus barbatus", "Aves", "Passeriformes", "Fringillidae"),
    ("Aguilucho", "Geranoaetus polyosoma", "Aves", "Accipitriformes", "Accipitridae"),
    ("Tórtola", "Zenaida auriculata", "Aves", "Columbiformes", "Columbidae"),
    ("Queltehue", "Vanellus chilensis", "Aves", "Charadriiformes", "Charadriidae"),
    ("Lagartija lemniscata", "Liolaemus lemniscatus", "Reptiles", "Squamata", "Liolaemidae"),
    ("Lagarto de Zapallar", "Liolaemus zapallarensis", "Reptiles", "Squamata", "Liolaemidae"),
    ("Lagartija oscura", "Liolaemus fuscus", "Reptiles", "Squamata", "Liolaemidae"),
    ("Lagarto nítido", "Liolaemus nitidus", "Reptiles", "Squamata", "Liolaemidae"),
    ("Sapito de cuatro ojos", "Pleurodema thaul", "Anfibios", "Anura", "Leptodactylidae"),
    ("Sapo de Atacama", "Rhinella atacamensis", "Anfibios", "Anura", "Bufonidae"),
    ("Sapo de rulo", "Rhinella arunco", "Anfibios", "Anura", "Bufonidae"),
    ("Zorro culpeo", "Lycalopex culpaeus", "Mamíferos", "Carnivora", "Canidae"),
    ("Cururo", "Spalacopus cyanus", "Mamíferos", "Rodentia", "Octodontidae"),
    ("Degú", "Octodon degus", "Mammalia", "Rodentia", "Octodontidae"),  # clase mal escrita (B01)
    ("Ratón orejudo de Darwin", "Phyllotis darwini", "Mamíferos", "Rodentia", "Cricetidae"),
    ("Conejo", "Oryctolagus cuniculus", "Mamíferos", "Lagomorpha", "Leporidae"),
    ("Perro", "Canis familiaris", "Mamíferos", "Carnivora", "Canidae"),
    ("Ratón sp", "-", "Mamíferos", "Rodentia", "-"),
    ("Sapo de pecho espinoso ", "Alsodes nodosus", "Anfibios", "Anura", "Alsodidae"),  # espacio sobrante (B01)
]

LISTAS = {
    "A": ("Metodología usada para registro", ["Cámara trampa", "Playback anfibios", "Playback aves",
                                               "Trampas Sherman", "Transecto", None, "Songmeter o similar", "Otro"]),
    "B": ("Comportamiento (añadir más si es necesario, dejar constancia)", [
        "Alimentándose", "Corriendo", "Cortejo", "Durmiendo", "En suelo", "Muerto", "Posado", "Tomando agua",
        "Volando", "Cazando", "Domesticado", "En competencia", "Vocalizando", "En cortejo",
        "En época reproductiva", "En vigilia", "Juego", "Reposando", "Otro", "Registro de audio"]),
    "C": ("Categoría de altura", [1, 2, 3, 4, 5]),
    "D": ("dirección vuelo", ["N", "S", "E", "O", "NE", "SE", "NO", "SO", "Otro"]),
    "E": ("Ladera", ["Norte", "Sur", "Este", "Oeste", "Noreste", "Sureste", "Noroeste", "Suroeste", "Otro", "Plano"]),
    "F": ("Tipo de registro", ["Individuo", "Vocalización", "Fecas", "Madriguera", "Cururera", "Plumas", "Muda",
                               "Huesos", "Otro", "Huella", "Nido", "Registro de audio", "Egagrópila"]),
}

PROYECTO = "Quebrada Verde (ejemplo)"
SI = "Si"
# ID, ambiente, ladera, estación (E, S), inicio, fin, metodologías (Q..U)
ESTACIONES = [
    ("QV01", "Matorral esclerófilo", "Norte", (300100, 6350100), (300030, 6350030), (300170, 6350170), (SI, 0, 0, 0, 0)),
    ("QV02", "Matorral esclerófilo", "Sur", (300500, 6350150), (300400, 6350150), (300600, 6350150), (SI, 0, 0, 0, 0)),
    ("QV03", "Matorral Esclerófilo", "Este", (300900, 6350200), (300900, 6350100), (300900, 6350300), (1, 0, 0, 0, 0)),
    ("QV04", "Esclerófilo", "Oeste", (301300, 6350250), (301230, 6350180), (301370, 6350320), (1, 0, 0, 0, 1)),
    ("QV05", "Bosque nativo", "Plano", (NBSP + "301700", 6350300), (301600, 6350300), (301800, 6350300), (1, 1, 1, 0, 0)),
    ("QV06", None, "Norte", (302100, 6350350), (302070, 6350350), (302130, 6350350), (1, 0, 0, 0, 0)),
    ("QV07", "Bosque nativo", "Sur", (300550, 6350250), (300450, 6350330), (300600, 6350150), ("SI", 0, 0, 0, 0)),
    ("QV08", "Matorral esclerófilo", "Este", (302500, 6350400), (302400, 6350400), (302600, 6350400), (1, 0, 0, 0, 0)),
    ("QV09", "Pradera", "Plano", (302900, 6355450), (302800, 6355450), (303000, 6355450), (1, 0, 0, 0, 0)),
]

D = date(2026, 10, 14)
# estación, día, hora, metodología, tipo, nombre común, abundancia, comportamiento, fotos, observaciones, captura (E, S)
REGISTROS = [
    ("QV01", 0, time(7, 40), "Transecto", "Vocalización", "Chincol", 2, None, None, None, None),
    ("QV01", 0, time(7, 45), "Transecto", "Individuo", "Loica", 1, "Posado", None, None, None),
    ("QV01", 0, time(7, 52), "Transecto", "Individuo", "Cometocino de gay", 3, "Alimentándose", None, None, None),
    ("QV01", 0, time(7, 52), "Transecto", "Individuo", "Cometocino de gay", 3, "Alimentándose", None, None, None),
    ("QV01", 0, time(8, 5), "Transecto", "Individuo", "lagartija lemniscata", 1, "Escondido", None, None, (300110, 6350090)),
    ("QV02", 0, time(8, 30), "Transecto", "Vocalización", "Tenca", "varios", None, None, None, None),
    ("QV02", 0, time(8, 41), "Transecto", "Individuo", "Lagartija lemniscata", 1, "Posado", "TC_00012-TC_00013", None, (300480, 6350160)),
    ("QV02", 0, time(8, 50), "Transecto", "Vocalización", "Cururo", "P", None, None, None, None),
    ("QV02", 0, time(9, 2), "Transecto", "Individuo", "Perro", 2, "Corriendo", None, None, None),
    ("QV03", 1, time(12, 15), "Transecto", "Vocalización", "Chincol", 1, None, None, None, None),
    ("QV03", 1, time(12, 40), "Transecto", "Individuo", "Diuca", 4, "Volando", None, None, None),
    ("QV03", 1, time(13, 5), "Transecto", "Individuo", "Rhinella spp.", 1, "Escondido", None, "La foto la tiene Ana, creo", (300905, 6350210)),
    ("QV03", 1, time(13, 20), "Cámara trampa", "Individuo", "Zorro culpeo", 1, "Corriendo", "TimePhoto_20261016_132000", None, None),
    ("QV03", 1, None, "Transecto", "Vocalización", "Turca", 1, None, None, None, None),
    ("QV04", 1, time(7, 10), "Trampas Sherman", "Individuo", "Ratón orejudo de Darwin", 1, None, None, None, (301240, 6350190)),
    ("QV04", 1, time(7, 25), "Trampas Sherman", "Individuo", "Chincol", 1, None, None, None, (301250, 6350200)),
    ("QV04", 1, time(14, 30), "Transecto", "Individuo", "Lagarto de zapallar", 1, "Posado", "TC_00021, TC_00022", None, (301300, 6350260)),
    ("QV04", 1, time(14, 45), "Transecto", "Huella", "Zorro culpeo", 1, None, None, None, None),
    ("QV05", 2, time(13, 30), "Transecto", "Vocalización", "Picaflor chico", 1, None, None, None, None),
    ("QV05", 2, time(14, 0), "Playback anfibios", "Vocalización", "Sapito de cuatro ojos", 3, None, None, None, None),
    ("QV05", 2, time(14, 10), "Transecto", "Individuo", "Lagartija oscura", 1, "Posado", None, "En sherman", (301650, NBSP + "6350310")),
    ("QV05", 2, time(15, 0), "Transecto", "Vocalización", "Perdiz chilena", 2, None, None, None, None),
    ("QV05", 2, time(15, 20), "transecto", "Individuo", "Jilguero", 6, "Volando", None, None, None),
    ("QV06", 2, time(12, 5), "Transecto", "Individuo", "Chincol", 3, "Posado", None, None, None),
    ("QV06", 2, time(12, 30), "Transecto", "Vocalización", "Loica", 1, None, None, None, None),
    ("QV06", 2, time(13, 10), "Transecto", "Individuo", "Lagarto nítido", 1, "Posado", None, None, (303500, 6352000)),
    ("QV07", 3, time(12, 45), "Transecto", "Vocalización", "Chercán", 1, None, None, None, None),
    ("QV07", 3, time(13, 0), "Transecto", "Individuo", "Tórtola", 2, "Volando", None, None, None),
    ("QV07", 3, time(13, 30), "Transecto", "Vocalización", "Chincol", 1, None, None, None, None),
    ("QV09", 3, time(16, 0), "Transecto", "Individuo", "Queltehue", 2, "En suelo", None, None, None),
    ("QV09", 3, time(16, 20), "Transecto", "Individuo", "Degú", 4, "Corriendo", None, None, None),
    ("QV09", 3, time(16, 40), "Transecto", "Vocalización", "Aguilucho", 1, None, None, "¿juvenil?", None),
    ("QV99", 3, time(10, 0), "Transecto", "Vocalización", "Chincol", 1, None, None, None, None),
    ("QV01", 21, time(9, 0), "Transecto", "Individuo", "Conejo", 1, "Corriendo", None, None, None),
    ("QV02", 3, time(10, 30), "Transecto", "Individuo", "Chincol", 40, "Volando", None, None, None),
    ("QV02", 3, time(10, 40), "Transecto", "Individuo", "Chincol", 2, "Volando", None, None, None),
    ("QV07", 3, time(11, 0), "Transecto", "Individuo", "Chincol", 1, "Posado", None, None, None),
    ("QV09", 3, time(17, 0), "Transecto", "Individuo", "Chincol", 3, "En suelo", None, None, None),
    ("QV03", 1, time(12, 50), "Transecto", "Individuo", "Chincol", 2, "En suelo", None, None, None),
]


def _formulas_registro(r: int) -> dict[str, str]:
    f = {
        letra: f"=INDEX(Estaciones!{col}$3:{col}$1000,MATCH($L{r},Estaciones!$H$3:$H$1000,0))"
        for letra, col in DESDE_ESTACIONES.items()
    }
    f.update({
        letra: f"=INDEX('nombres BD'!{col}:{col},MATCH($AD{r},'nombres BD'!$A:$A,0))" for letra, col in DESDE_BD.items()
    })
    f.update({
        "E": f"=YEAR(D{r})", "F": f"=MONTH(D{r})", "G": f"=DAY(D{r})",
        "Y": f'=IF(HOUR(Z{r})>11,"PM",IF(HOUR(Z{r})<12,"AM","REVISAR"))',
        "AB": "=" + FORMULA_DIR_IND.format(r=r),
    })
    return f


def crear_planilla_demo(destino: str | Path) -> Path:
    """Escribe la planilla de ejemplo en ``destino`` y devuelve la ruta."""
    destino = Path(destino)
    wb = openpyxl.Workbook()
    leer = wb.active
    leer.title = "LEER ANTES DE USAR"
    leer["A1"] = "Planilla de EJEMPLO con datos ficticios para probar el Auditor XPE."
    leer["A2"] = "Tiene errores a propósito: coordenadas con espacios invisibles, nombres mal escritos, duplicados…"

    est = wb.create_sheet("Estaciones")
    for letra, texto in (("A", "Descripción estación"), ("H", "ID"), ("K", "Coordenadas estación, transectos y Tránsito Aéreo"),
                         ("Q", "Metodologías usadas")):
        est[f"{letra}1"] = texto
    for j, enc in enumerate(ENC_ESTACIONES, start=1):
        c = est.cell(2, j, enc)
        c.font = NEGRITA
    for i, (id_, amb, ladera, p, ini, fin, mets) in enumerate(ESTACIONES, start=1):
        r = i + 2
        proyecto = PROYECTO if id_ != "QV06" else PROYECTO[:-3]  # variante truncada (E13)
        fila = [i, proyecto, "Valparaíso", "Primavera", None, amb, ladera, id_, None if id_ == "QV08" else id_, None,
                *p, *ini, *fin, *mets, "Persona Ejemplo", "Persona Ejemplo"]
        for j, v in enumerate(fila, start=1):
            est.cell(r, j, v)
    dv = DataValidation(type="list", formula1="'validación datos'!$E$2:$E$11", allow_blank=True)
    dv.add("G3:G100")
    est.add_data_validation(dv)

    reg = wb.create_sheet("Registros")
    reg["B1"] = "XPE CONSULT"
    reg["B2"] = "FORMULARIO FAUNA TERRESTRE (EJEMPLO)"
    for j, enc in enumerate(ENC_REGISTROS, start=1):
        c = reg.cell(9, j, enc)
        c.font = NEGRITA
    letras = {openpyxl.utils.get_column_letter(j): j for j in range(1, len(ENC_REGISTROS) + 1)}
    for r in range(10, 10 + FILAS_PLANTILLA):
        for letra, formula in _formulas_registro(r).items():
            reg.cell(r, letras[letra], formula).fill = NARANJO
    for i, (id_, dia, hora, met, tipo, nombre, ab, comp, fotos, obs, cap) in enumerate(REGISTROS):
        r = 10 + i
        valores = {
            "A": i + 1, "D": datetime.combine(D + timedelta(days=dia), time()),
            "K": "Despejado", "L": id_, "Z": hora, "AA": met, "AC": tipo, "AD": nombre, "AF": "Adulto",
            "AG": "Vivo", "AH": "Indeterminado", "AI": ab, "AK": fotos, "AL": obs,
        }
        if comp is not None:
            valores["AJ"] = comp
        elif tipo == "Vocalización":
            valores["AJ"] = f'=IF(AC{r}="Vocalización","Vocalización","COMPLETAR")'  # fórmula de ayuda
        if cap:
            valores["AM"], valores["AN"] = cap
        for letra, v in valores.items():
            if v is not None:
                c = reg.cell(r, letras[letra], v)
                if letra == "D":
                    c.number_format = "dd-mm-yyyy"
                if letra == "Z":
                    c.number_format = "hh:mm"
    for rango, formula in (("AA10:AA200", "'validación datos'!$A$2:$A$9"), ("AC10:AC200", "'validación datos'!$F$2:$F$14")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rango)
        reg.add_data_validation(dv)

    ta = wb.create_sheet("Registros TA")
    for j, enc in enumerate(ENC_TA, start=1):
        ta.cell(9, j, enc).font = NEGRITA

    bd = wb.create_sheet("nombres BD")
    for j, enc in enumerate(["Nombre común (1)", "Nombre científico (2)", "Reino (3)", "Tipo (4)", "Clase (5)",
                             "Orden (6)", "Familia (7)", "Género", "Epíteto específico", "sub especie"], start=1):
        bd.cell(1, j, enc).font = NEGRITA
    for i, (comun, cientifico, clase, orden, familia) in enumerate(ESPECIES, start=2):
        partes = cientifico.split(" ")
        genero, epiteto = (partes + ["-"])[:2] if cientifico != "-" else ("-", "-")
        for j, v in enumerate([comun, cientifico, "Animalia", "Chordata", clase, orden, familia, genero, epiteto, "-"], 1):
            bd.cell(i, j, v)

    val = wb.create_sheet("validación datos")
    for letra, (enc, valores) in LISTAS.items():
        val[f"{letra}1"] = enc
        for i, v in enumerate(valores, start=2):
            if v is not None:
                val[f"{letra}{i}"] = v
    wb.active = 1
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
