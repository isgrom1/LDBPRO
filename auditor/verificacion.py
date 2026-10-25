"""Verificaciones del archivo _CORREGIDA.xlsx contra el original."""

from __future__ import annotations

import tempfile
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from . import xlsx_xml
from .formato import convertir_con_libreoffice, soffice_disponible

ERRORES_EXCEL = {"#N/A", "#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


@dataclass
class Verificacion:
    x14_original: dict[str, int]
    x14_corregida: dict[str, int]
    formulas_original: dict[str, int]
    formulas_corregida: dict[str, int]
    estilos_identicos: bool
    recalculo_al_abrir: bool
    celdas_distintas: set[tuple[str, str]]
    celdas_esperadas: set[tuple[str, str]]
    libreoffice: dict | None = None
    reauditoria_coincide: bool | None = None
    diferencias_reauditoria: list[str] = field(default_factory=list)

    @property
    def x14_ok(self) -> bool:
        return self.x14_original == self.x14_corregida

    @property
    def formulas_ok(self) -> bool:
        return self.formulas_original == self.formulas_corregida

    @property
    def celdas_ok(self) -> bool:
        return self.celdas_distintas == self.celdas_esperadas

    @property
    def libreoffice_ok(self) -> bool | None:
        if self.libreoffice is None:
            return None
        return not self.libreoffice["errores_nuevos"] and not self.libreoffice.get("diferencias_uniones")

    @property
    def todo_ok(self) -> bool:
        checks = [self.x14_ok, self.formulas_ok, self.estilos_identicos, self.recalculo_al_abrir, self.celdas_ok]
        if self.libreoffice_ok is not None:
            checks.append(self.libreoffice_ok)
        if self.reauditoria_coincide is not None:
            checks.append(self.reauditoria_coincide)
        return all(checks)


def _celdas(ruta: Path) -> dict[str, dict[str, tuple]]:
    """{hoja: {ref: (valor, es_fórmula)}} con valores crudos (no cacheados)."""
    out = {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=False)
        try:
            for ws in wb.worksheets:
                d = {}
                for fila in ws.iter_rows():
                    for c in fila:
                        if getattr(c, "value", None) is not None:
                            d[c.coordinate] = (c.value, c.data_type == "f")
                out[ws.title] = d
        finally:
            wb.close()
    return out


def errores_de_formulas(ruta: Path) -> dict[tuple[str, str], str]:
    """Celdas con fórmula cuyo valor (cacheado) es un error de Excel."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wf = openpyxl.load_workbook(ruta, read_only=True, data_only=False)
        wv = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    out = {}
    try:
        for hoja in wf.sheetnames:
            for ff, fv in zip(wf[hoja].iter_rows(), wv[hoja].iter_rows()):
                for cf, cv in zip(ff, fv):
                    if getattr(cf, "data_type", None) == "f" and isinstance(cv.value, str) and (
                        cv.value in ERRORES_EXCEL or cv.value.startswith("Err:")
                    ):
                        out[(hoja, cf.coordinate)] = cv.value
    finally:
        wf.close()
        wv.close()
    return out


def verificar_corregida(original: Path, corregida: Path, cambios: dict[str, dict[str, object]],
                        con_libreoffice: bool = True, config: dict | None = None) -> Verificacion:
    with zipfile.ZipFile(original) as zo, zipfile.ZipFile(corregida) as zc:
        estilos = zo.read("xl/styles.xml") == zc.read("xl/styles.xml")
        recalculo = 'fullCalcOnLoad="1"' in zc.read("xl/workbook.xml").decode("utf-8")
    a, b = _celdas(original), _celdas(corregida)
    distintas = set()
    for hoja in set(a) | set(b):
        da, db = a.get(hoja, {}), b.get(hoja, {})
        for ref in set(da) | set(db):
            if da.get(ref) != db.get(ref):
                distintas.add((hoja, ref))
    # Un cambio a "vacío" sobre una celda que ya estaba vacía no deja diferencia visible.
    esperadas = {(h, r) for h, cs in cambios.items() for r, v in cs.items() if a.get(h, {}).get(r, (None,))[0] != v}
    v = Verificacion(
        xlsx_xml.contar_x14_validaciones(original), xlsx_xml.contar_x14_validaciones(corregida),
        xlsx_xml.contar_formulas(original), xlsx_xml.contar_formulas(corregida),
        estilos, recalculo, distintas, esperadas,
    )
    if con_libreoffice and soffice_disponible():
        v.libreoffice = recalcular_y_comparar(original, corregida, config)
    return v


def recalcular_y_comparar(original: Path, corregida: Path, config: dict | None) -> dict:
    """Recalcula ambos archivos con LibreOffice y compara los errores de las fórmulas.

    Además vuelve a correr X02 (caché vs. unión rehecha en código) sobre la corregida
    recalculada: si las uniones del auditor emulan bien las fórmulas, no hay diferencias.
    """
    with tempfile.TemporaryDirectory(prefix="verif_lo_") as tmp:
        ro = convertir_con_libreoffice(original, Path(tmp) / "original")
        rc = convertir_con_libreoffice(corregida, Path(tmp) / "corregida")
        eo, ec = errores_de_formulas(ro), errores_de_formulas(rc)
        nuevos = sorted(set(ec) - set(eo))
        resueltos = sorted(set(eo) - set(ec))
        diferencias = None
        if config is not None:
            from .contexto import Contexto
            from .modelo import leer_libro
            from .reglas.plantilla import x02

            libro = leer_libro(rc, config)
            diferencias = [f"{h.hoja}!{h.celda}: {h.problema}" for h in x02(Contexto(libro, config))]
    return {
        "errores_original": len(eo),
        "errores_corregida": len(ec),
        "errores_nuevos": [f"{h}!{r} = {ec[(h, r)]}" for h, r in nuevos],
        "errores_resueltos": [f"{h}!{r}" for h, r in resueltos],
        "diferencias_uniones": diferencias,
    }
