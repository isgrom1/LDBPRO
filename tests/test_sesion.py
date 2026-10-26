"""Decisiones del usuario sobre la planilla de ejemplo (no necesita la planilla real)."""

from __future__ import annotations

import csv
import warnings
import zipfile

import openpyxl
import pytest

from auditor.config import cargar_config
from auditor.demo import crear_planilla_demo
from auditor.sesion import (
    ErrorDecision,
    analizar,
    cambios_con_decisiones,
    convertir_valor,
    generar,
    preparar_decision,
    proyectar,
)


@pytest.fixture(scope="module")
def analisis(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("demo")
    ruta = crear_planilla_demo(carpeta / "planilla_ejemplo.xlsx")
    return analizar(ruta, cargar_config(), carpeta, informar=lambda *a: None)


def buscar(analisis, regla, celda=None):
    return next(h for h in analisis.antes if h.id_regla == regla and (celda is None or h.celda == celda))


def test_la_demo_muestra_errores_de_todas_las_familias(analisis):
    reglas = {h.id_regla for h in analisis.antes}
    esperadas = {"E02", "E03", "E04", "E05", "E07", "E08", "E09", "E11", "E12", "E13", "R01", "R02", "R03",
                 "R04", "R05", "R06", "R08", "R09", "R10", "R12", "R13", "R14", "R15", "R16", "R17", "R18",
                 "R19", "R20", "R21", "R22", "R23", "R25", "R26", "B01", "T01"}
    assert esperadas <= reglas, esperadas - reglas
    assert not [h for h in analisis.antes if h.id_regla == "ERROR"]
    assert analisis.puntaje_despues.indice > analisis.puntaje_antes.indice


def test_horarios_y_origen(analisis):
    # R12 solo marca el playback hecho claramente de día (QV05: Playback anfibios a las 14:00)
    assert [h.id_estacion for h in analisis.antes if h.id_regla == "R12"] == ["QV05"]
    conejo = buscar(analisis, "R25")
    assert conejo.valor_actual == "Conejo" and conejo.id_estacion == "QV01"


def test_presencia_duplicados_y_motivos_de_terreno(analisis, tmp_path):
    # «P» es presencia: aviso, no error; «varios» sí es error y ofrece P o 1
    r26 = buscar(analisis, "R26")
    assert r26.severidad.nivel == 1 and r26.valor_actual == "P" and r26.correccion is None
    assert convertir_valor("abundancia", " presencia ") == "P"
    # filas idénticas: advertencia, con botón para marcar el duplicado en Observaciones
    r15 = buscar(analisis, "R15")
    assert r15.severidad.nivel == 2 and r15.justificaciones == ["Son registros distintos"]
    assert r15.acciones[0].etiqueta == "Es el mismo: marcar como duplicado"
    primera = int(r15.problema.split("Fila idéntica a la ")[1].split(":")[0])
    assert r15.acciones[0].valor == f"[DUPLICADO de la fila {primera}]"
    # 40 chincoles: paseriforme, ave de bandada → R07 no lo marca; sin esa excepción, sí
    assert not [h for h in analisis.antes if h.id_regla == "R07"]
    cfg = cargar_config()
    cfg["abundancia"]["ordenes_gregarios"] = []
    sin_excepcion = analizar(analisis.ruta, cfg, tmp_path / "sin_excepcion", informar=lambda *a: None)
    r07 = next(h for h in sin_excepcion.antes if h.id_regla == "R07")
    assert r07.valor_actual == 40 and r07.justificaciones == ["Conteo real"]

    decisiones = {r15.clave(): preparar_decision(analisis, r15, {"tipo": "accion", "indice": 0})}
    r = generar(analisis, tmp_path, decisiones, con_libreoffice=False, informar=lambda *a: None)
    assert r.verificacion.todo_ok
    assert not [h for h in r.hallazgos_despues if h.id_regla == "R15" and h.celda == r15.celda]
    filas = list(csv.DictReader(r.archivos["base"].read_text(encoding="utf-8").splitlines()))
    cururo = next(f for f in filas if f["nombre_comun"] == "Cururo")
    assert (cururo["abundancia"], cururo["tipo_abundancia"], cururo["usar_en_abundancia"]) == ("", "presencia", "no")
    assert cururo["excluir_motivo"] == ""  # entra a la riqueza
    duplicada = next(f for f in filas if f["fila"] == str(r15.fila))
    assert duplicada["excluir_motivo"] == f"duplicado de la fila {primera}"
    assert "Registros de presencia («P»" in r.archivos["reporte"].read_text(encoding="utf-8")


def test_nombre_comun_con_dos_especies(tmp_path):
    ruta = crear_planilla_demo(tmp_path / "p.xlsx")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(ruta)
    bd = wb["nombres BD"]
    bd.append(["Chincol", "Zonotrichia otra", "Animalia", "Chordata", "Aves", "Passeriformes", "Passerellidae", "Zonotrichia"])
    wb.save(ruta)
    a = analizar(ruta, cargar_config(), tmp_path, informar=lambda *x: None)
    b01 = [h for h in a.antes if h.id_regla == "B01" and "Chincol" in h.problema]
    assert len(b01) == 1 and "completa la primera (Zonotrichia capensis)" in b01[0].problema
    r27 = [h for h in a.antes if h.id_regla == "R27"]
    assert r27 and all("Zonotrichia otra" in h.problema for h in r27)
    assert r27[0].justificaciones == ["Es Zonotrichia capensis"]


def test_acciones_sugeridas(analisis):
    assert [a.etiqueta for a in buscar(analisis, "R06").acciones] == ["Poner P (presencia)", "Poner 1"]
    assert buscar(analisis, "E03", "F6").acciones[0].valor == "Matorral esclerófilo"
    r20 = buscar(analisis, "R20")
    assert (r20.acciones[0].hoja, r20.acciones[0].valor) == ("Estaciones", 1)
    assert "Sapo de Atacama" in [a.valor for a in buscar(analisis, "R02").acciones]
    assert buscar(analisis, "R13").opciones[0] == "Cámara trampa"


def test_decisiones_validas_e_invalidas(analisis):
    r06 = buscar(analisis, "R06")
    d = preparar_decision(analisis, r06, {"tipo": "valor", "texto": " 2 ", "nota": "contados en terreno"})
    assert (d.valor, d.celda, d.nota) == (2, r06.celda, "contados en terreno")
    with pytest.raises(ErrorDecision, match="no es un número"):
        preparar_decision(analisis, r06, {"tipo": "valor", "texto": "muchos"})
    with pytest.raises(ErrorDecision, match="fórmula"):
        preparar_decision(analisis, r06, {"tipo": "valor", "texto": "x", "hoja": "Registros", "celda": "AE10"})
    with pytest.raises(ErrorDecision, match="no es una celda de datos"):
        preparar_decision(analisis, r06, {"tipo": "valor", "texto": "1", "hoja": "Registros", "celda": "AI1"})
    fecha = preparar_decision(analisis, buscar(analisis, "R09"), {"tipo": "valor", "texto": "16-10-2026"})
    assert fecha.valor == 46311


def test_decisiones_cambian_los_cambios(analisis):
    r03 = buscar(analisis, "R03", "AD12")
    r20 = buscar(analisis, "R20")
    decisiones = {
        r03.clave(): preparar_decision(analisis, r03, {"tipo": "sin_auto"}),
        r20.clave(): preparar_decision(analisis, r20, {"tipo": "accion", "indice": 0}),
        buscar(analisis, "R06").clave(): preparar_decision(analisis, buscar(analisis, "R06"), {"tipo": "accion", "indice": 1}),
    }
    cambios, resueltos = cambios_con_decisiones(analisis, decisiones)
    celdas = {(c.hoja, c.celda): c for c in cambios}
    assert ("Registros", "AD12") not in celdas  # auto-corrección rechazada
    assert ("Registros", "AD13") in celdas  # la otra sigue
    assert celdas[(r20.acciones[0].hoja, r20.acciones[0].celda)].valor_nuevo == 1
    assert celdas[("Registros", buscar(analisis, "R06").celda)].reglas == ["MANUAL:R06"]
    assert r03.clave() not in resueltos and r20.clave() in resueltos


def test_generar_con_decisiones(analisis, tmp_path):
    r14 = buscar(analisis, "R14")  # Directo/Indirecto: celda con fórmula compartida que se puede sobrescribir
    r15 = buscar(analisis, "R15")
    r06 = buscar(analisis, "R06")
    decisiones = {
        r14.clave(): preparar_decision(analisis, r14, {"tipo": "accion", "indice": 1}),
        r15.clave(): preparar_decision(analisis, r15, {"tipo": "ok", "nota": "son dos individuos"}),
        r06.clave(): preparar_decision(analisis, r06, {"tipo": "valor", "texto": "1"}),
    }
    assert decisiones[r14.clave()].reemplaza_formula
    _, proyectado = proyectar(analisis, decisiones)
    r = generar(analisis, tmp_path, decisiones, con_libreoffice=False, informar=lambda *a: None)
    assert r.verificacion.todo_ok, r.verificacion
    assert r.puntaje_despues.indice == pytest.approx(proyectado.indice)
    assert r.aceptados == 1
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(r.archivos["corregida"])
        rev = openpyxl.load_workbook(r.archivos["revision"])
    reg = wb["Registros"]
    fila = int(r14.celda[2:])
    assert reg[r14.celda].value == "Indirecto"
    assert str(reg[f"AB{fila + 1}"].value).startswith("=IF(OR(AC")  # el resto de la columna sigue con fórmula
    assert reg[r06.celda].value == 1
    h = rev["HALLAZGOS"]
    columnas = [c.value for c in h[1]]
    decisiones_escritas = [h.cell(i, columnas.index("Decisión / Corrección") + 1).value for i in range(2, h.max_row + 1)]
    assert any(d and "son dos individuos" in d for d in decisiones_escritas)
    cambios = list(csv.DictReader(open(r.archivos["cambios"], encoding="utf-8")))
    assert {c["regla"] for c in cambios} >= {"MANUAL:R14", "MANUAL:R06"}
    assert "## Decisiones tomadas en la revisión (3)" in r.archivos["reporte"].read_text(encoding="utf-8")
    with zipfile.ZipFile(r.archivos["corregida"]) as z:
        assert "xl/calcChain.xml" not in z.namelist()
