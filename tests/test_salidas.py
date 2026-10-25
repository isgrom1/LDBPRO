"""Carpeta de salida, archivo de revisión, CSV y CLI."""

from __future__ import annotations

import csv
import shutil
import subprocess
import sys
import warnings
from pathlib import Path

import openpyxl
import pytest

from auditor.cli import ejecutar
from auditor.formato import (
    convertir_con_libreoffice,
    detectar_formato,
    soffice_disponible,
)
from auditor.reportes import COLUMNAS_BASE

RAIZ = Path(__file__).resolve().parent.parent


def test_carpeta_y_archivos(resultado):
    assert resultado.carpeta.name == "auditoria_Formato_Planilla_Fauna_XPE_Externo_v4_version_1_2026-10-24"
    nombres = sorted(p.name for p in resultado.carpeta.iterdir())
    assert nombres == sorted([
        "Formato_Planilla_Fauna_XPE_Externo_v4_version_1_REVISION.xlsx",
        "Formato_Planilla_Fauna_XPE_Externo_v4_version_1_CORREGIDA.xlsx",
        "cambios_aplicados.csv", "reporte.md", "base_limpia_larga.csv", "hallazgos.csv",
    ])


def test_no_reutiliza_una_carpeta_existente(planilla, tmp_path):
    a = ejecutar(planilla, salida=tmp_path, con_libreoffice=False, informar=lambda *x: None, solo_reporte=True)
    b = ejecutar(planilla, salida=tmp_path, con_libreoffice=False, informar=lambda *x: None, solo_reporte=True)
    assert a.carpeta != b.carpeta and b.carpeta.name.endswith("_2")


def test_solo_reporte_no_escribe_corregida(planilla, tmp_path):
    r = ejecutar(planilla, salida=tmp_path, con_libreoffice=False, informar=lambda *x: None, solo_reporte=True)
    nombres = {p.name for p in r.carpeta.iterdir()}
    assert not any("CORREGIDA" in n for n in nombres) and "cambios_aplicados.csv" not in nombres
    assert "Proyectado" in (r.carpeta / "reporte.md").read_text(encoding="utf-8")


def test_archivo_de_revision(resultado):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(resultado.archivos["revision"])
    assert wb.sheetnames[:2] == ["RESUMEN", "HALLAZGOS"]
    h = wb["HALLAZGOS"]
    encabezados = [c.value for c in h[1]]
    for col in ("Regla", "Hoja", "Celda", "Fila", "ID Estación", "Valor actual", "Problema", "Sugerencia",
                "Severidad", "Auto-corregible", "Decisión / Corrección"):
        assert col in encabezados
    assert h.auto_filter.ref == f"A1:P{len(resultado.hallazgos_antes) + 1}"
    assert all(h.cell(i, encabezados.index("Decisión / Corrección") + 1).value is None for i in range(2, h.max_row + 1))
    rojo, naranjo, amarillo = "FFFF7C80", "FFFFC000", "FFFFFF66"
    assert wb["Estaciones"]["K19"].fill.fgColor.rgb == rojo  # E05
    assert wb["Estaciones"]["F15"].fill.fgColor.rgb == naranjo  # E03
    assert wb["Registros"]["AD198"].fill.fgColor.rgb == amarillo  # R08
    assert "R03" in wb["Registros"]["AD171"].comment.text
    resumen = " ".join(str(c.value) for fila in wb["RESUMEN"].iter_rows() for c in fila if c.value)
    for texto in ("Índice de calidad global", "% registros limpios", "% estaciones limpias", "Top 10", "❌ NO USAR"):
        assert texto in resumen


def test_cambios_aplicados_csv(resultado):
    filas = list(csv.DictReader(open(resultado.archivos["cambios"], encoding="utf-8")))
    assert len(filas) == len(resultado.cambios)
    assert set(filas[0]) >= {"hoja", "celda", "valor_anterior", "valor_nuevo", "regla"}
    k19 = next(f for f in filas if f["celda"] == "K19")
    assert k19["valor_anterior"] == " 251835" and k19["valor_nuevo"] == "251835" and k19["regla"] == "E05"


def test_base_limpia_larga(resultado):
    filas = list(csv.DictReader(open(resultado.archivos["base"], encoding="utf-8")))
    assert list(filas[0]) == COLUMNAS_BASE
    assert len(filas) == 198  # no se borran filas
    caballo = next(f for f in filas if f["nombre_cientifico"] == "Equus caballus")
    assert caballo["excluir_motivo"] == "doméstica" and caballo["usar_en_abundancia"] == "no"
    assert caballo["origen"] == "doméstica"
    rhinella = next(f for f in filas if f["nombre_comun"] == "Rhinella spp.")
    assert "nombres BD" in rhinella["excluir_motivo"]
    gay = [f for f in filas if f["nombre_cientifico"] == "Phrygilus gayi"]
    assert {f["nombre_comun"] for f in gay} == {"Cometocino de Gay"}  # unión rehecha con el nombre corregido
    pmf34 = next(f for f in filas if f["estacion"] == "PMF34")
    assert pmf34["ambiente"] == "Medio alto"  # la unificación de ambientes la decide el usuario (E03)
    assert next(f for f in filas if f["estacion"] == "PMF55")["largo_transecto_m"] == "71.6"


def test_reporte_md(resultado):
    md = resultado.archivos["reporte"].read_text(encoding="utf-8")
    for texto in ("## Puntaje de calidad", "## Qué se corrigió solo", "## Qué falta decidir", "### PMF34",
                  "## Verificación del archivo corregido", "## Mapeo de columnas detectado", "Fórmula:"):
        assert texto in md


def test_cli(planilla, tmp_path):
    proc = subprocess.run(
        [sys.executable, str(RAIZ / "auditar_planilla.py"), str(planilla), "--salida", str(tmp_path),
         "--sin-libreoffice", "--config", str(RAIZ / "config.yaml")],
        capture_output=True, text=True, encoding="utf-8", cwd=tmp_path,
    )
    assert proc.returncode == 0, proc.stderr
    assert "ANTES" in proc.stdout and "DESPUÉS" in proc.stdout and "Mapeo de columnas" in proc.stdout


def test_cli_archivo_inexistente(tmp_path):
    proc = subprocess.run([sys.executable, str(RAIZ / "auditar_planilla.py"), str(tmp_path / "no.xlsx")],
                          capture_output=True, text=True)
    assert proc.returncode == 1 and "No existe" in proc.stderr


@pytest.mark.libreoffice
@pytest.mark.skipif(not soffice_disponible(), reason="LibreOffice no está instalado")
def test_extension_enganosa_se_convierte(planilla, tmp_path):
    """Un .xls (binario) que llega con extensión .xlsx se detecta y se convierte antes de auditar."""
    xls_dir = tmp_path / "xls"
    xlsx = convertir_con_libreoffice(planilla, xls_dir, recalcular=False)
    subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", str(xls_dir), str(xlsx)],
                   capture_output=True, timeout=300, check=True)
    disfrazado = tmp_path / "planilla.xlsb.xlsx"
    shutil.copyfile(next(xls_dir.glob("*.xls")), disfrazado)
    assert detectar_formato(disfrazado) == "xls"
    r = ejecutar(disfrazado, salida=tmp_path, con_libreoffice=False, informar=lambda *x: None, solo_reporte=True)
    assert r.libro.formato == "xls" and len(r.libro.registros[0].filas_datos) == 198
    assert any(h.id_regla == "R03" for h in r.hallazgos_antes)
